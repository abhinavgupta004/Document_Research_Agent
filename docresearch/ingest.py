"""Loader -> Chunking -> Embeddings -> Vector DB.

Page numbers are carried through every step, because "page 7 of paper 2" is
what makes an answer checkable by a human.
"""
import hashlib
from dataclasses import dataclass, field
from pathlib import Path

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from .config import settings
from .vectorstore import add_chunks, delete_source

SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


@dataclass
class IngestReport:
    file: str
    chunks: int = 0
    pages: int = 0
    empty_pages: int = 0
    error: str | None = None
    notes: list[str] = field(default_factory=list)


# ---------------------------------------------------------------- loading
def load_file(path: Path) -> list[Document]:
    """Return one Document per page (PDF) or one per file (everything else)."""
    ext = path.suffix.lower()
    name = path.name

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [
            Document(
                page_content=(page.extract_text() or "").strip(),
                metadata={"source": name, "page": i + 1},
            )
            for i, page in enumerate(reader.pages)
        ]

    if ext in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
        return [Document(page_content=text, metadata={"source": name, "page": 1})]

    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        # Tables hold a lot of real content (tech stacks, requirements...), so keep them.
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                parts.append(" | ".join(dict.fromkeys(c for c in cells if c)))
        text = "\n".join(parts).strip()
        # Word has no fixed pages, so the whole file counts as page 1.
        return [Document(page_content=text, metadata={"source": name, "page": 1})]

    raise ValueError(f"Unsupported file type: {ext}")


# --------------------------------------------------------------- chunking
def chunk_documents(docs: list[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(docs)

    # Stable ids: re-uploading the same file updates instead of duplicating.
    counters: dict[tuple, int] = {}
    for c in chunks:
        key = (c.metadata["source"], c.metadata["page"])
        counters[key] = counters.get(key, 0) + 1
        raw = f"{key[0]}|{key[1]}|{counters[key]}"
        c.metadata["chunk_id"] = hashlib.md5(raw.encode()).hexdigest()
    return chunks


# -------------------------------------------------------------- ingestion
def ingest_file(path: str | Path) -> IngestReport:
    path = Path(path)
    report = IngestReport(file=path.name)

    if path.suffix.lower() not in SUPPORTED:
        report.error = f"{path.suffix or 'that'} files aren't supported yet."
        return report

    try:
        pages = load_file(path)
    except Exception as exc:  # corrupt / encrypted files happen
        report.error = f"Couldn't read this file ({exc})."
        return report

    report.pages = len(pages)
    non_empty = [p for p in pages if p.page_content]
    report.empty_pages = len(pages) - len(non_empty)

    if not non_empty:
        report.error = (
            "No text found. This looks like a scanned PDF - OCR isn't built in yet."
        )
        return report
    if report.empty_pages:
        report.notes.append(
            f"{report.empty_pages} page(s) had no extractable text (scans or images) and were skipped."
        )

    chunks = chunk_documents(non_empty)

    # If the file was indexed before, wipe the old version so stale chunks don't linger.
    delete_source(path.name)
    add_chunks(chunks, ids=[c.metadata["chunk_id"] for c in chunks])
    report.chunks = len(chunks)
    return report